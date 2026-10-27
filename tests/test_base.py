import os
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import date

import openpyxl

from database.db_manager import VERSION_SCHEMA, DBManager
from database.modeles import DEPENSE, RECETTE, ErreurMetier
from database.queries import TresoQueries
from services.bilan import Exercice
from services.export_excel import exporter_bilan

SCHEMA_V1 = """
    CREATE TABLE categories (id INTEGER PRIMARY KEY AUTOINCREMENT, nom TEXT NOT NULL UNIQUE,
        type_flux TEXT NOT NULL CHECK(type_flux IN ('RECETTE', 'DEPENSE')));
    CREATE TABLE tiers (id INTEGER PRIMARY KEY AUTOINCREMENT, nom TEXT NOT NULL UNIQUE,
        type_tier TEXT NOT NULL CHECK(type_tier IN ('PHYSIQUE', 'MORAL')));
    CREATE TABLE transactions (id INTEGER PRIMARY KEY AUTOINCREMENT, date_transaction TEXT NOT NULL,
        description TEXT NOT NULL, montant REAL NOT NULL CHECK(montant >= 0), reference_facture TEXT,
        categorie_id INTEGER NOT NULL, tier_id INTEGER,
        FOREIGN KEY (categorie_id) REFERENCES categories(id) ON DELETE RESTRICT,
        FOREIGN KEY (tier_id) REFERENCES tiers(id) ON DELETE SET NULL);
    INSERT INTO categories (nom, type_flux) VALUES ('Cotisations', 'RECETTE'), ('Boissons', 'DEPENSE');
    INSERT INTO tiers (nom, type_tier) VALUES ('Métro', 'MORAL');
    INSERT INTO transactions (date_transaction, description, montant, reference_facture, categorie_id, tier_id)
    VALUES ('2025-10-01', 'Adhésions', 630.0, NULL, 1, NULL), ('2025-10-02', 'Achat Métro', 19.99, 'F1', 2, 1);
"""


class BaseTemporaire(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.TemporaryDirectory()
        self.chemin = os.path.join(self.dossier.name, "tresorerie.db")

    def tearDown(self):
        self.dossier.cleanup()


class TestSchema(BaseTemporaire):
    def test_base_neuve_avec_plan_de_depart(self):
        q = TresoQueries(DBManager(self.chemin))
        self.assertEqual([c.code for c in q.get_tous_comptes()], ["512", "507", "530 001", "530 002", "530 003"])
        adhesions = next(c for c in q.get_toutes_categories() if c.nom == "Adhésions")
        self.assertIsNotNone(adhesions.parent_id)
        self.assertEqual(q.get_parametre("nom_association"), "BDE ODIN")

    def test_migration_depuis_la_version_1(self):
        conn = sqlite3.connect(self.chemin)
        conn.executescript(SCHEMA_V1)
        conn.close()

        q = TresoQueries(DBManager(self.chemin))
        operations = sorted(q.get_toutes_operations(), key=lambda op: op.id)
        self.assertEqual([op.montant for op in operations], [63000, 1999])
        self.assertEqual({op.compte_nom for op in operations}, {"Compte courant"})
        self.assertEqual(operations[1].tier_nom, "Métro")
        with closing(sqlite3.connect(self.chemin)) as conn:
            self.assertEqual(conn.execute("PRAGMA user_version").fetchone()[0], VERSION_SCHEMA)
            self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
        DBManager(self.chemin)  # une seconde ouverture ne refait pas la migration

    def test_import_d_une_base_v1(self):
        ancienne = os.path.join(self.dossier.name, "v1.db")
        with closing(sqlite3.connect(ancienne)) as conn:
            conn.executescript(SCHEMA_V1)
        db = DBManager(self.chemin)
        sauvegarde = db.importer(ancienne)
        self.assertTrue(os.path.exists(sauvegarde))
        self.assertEqual(len(TresoQueries(db).get_toutes_operations()), 2)
        with closing(sqlite3.connect(sauvegarde)) as conn:  # l'ancienne base 2.x reste intacte
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM comptes").fetchone()[0], 5)

    def test_import_refuse_un_fichier_quelconque(self):
        autre = os.path.join(self.dossier.name, "notes.db")
        with open(autre, "w", encoding="utf-8") as f:
            f.write("pas une base")
        with self.assertRaisesRegex(ErreurMetier, "pas une base OdinFinance"):
            DBManager(self.chemin).importer(autre)


class TestRegles(BaseTemporaire):
    def setUp(self):
        super().setUp()
        self.q = TresoQueries(DBManager(self.chemin))
        self.courant = self.q.get_tous_comptes()[0].id
        self.livret = self.q.get_tous_comptes()[1].id
        self.ventes = next(c.id for c in self.q.get_toutes_categories() if c.nom == "Ventes")

    def test_doublon_signale_lisiblement(self):
        with self.assertRaisesRegex(ErreurMetier, "existe déjà"):
            self.q.ajouter_categorie("Ventes", RECETTE)

    def test_sous_categorie_un_seul_niveau_et_meme_type(self):
        adhesions = next(c.id for c in self.q.get_toutes_categories() if c.nom == "Adhésions")
        with self.assertRaises(ErreurMetier):
            self.q.ajouter_categorie("Cotisation annuelle", RECETTE, adhesions)
        with self.assertRaises(ErreurMetier):
            self.q.ajouter_categorie("Matériel", DEPENSE, self.ventes)

    def test_suppressions_protegees(self):
        self.q.enregistrer_operation(None, date(2025, 5, 1), "Vente", 1000, None, self.courant, self.ventes, None, None)
        with self.assertRaisesRegex(ErreurMetier, "1 opération"):
            self.q.supprimer_categorie(self.ventes)
        with self.assertRaisesRegex(ErreurMetier, "opération"):
            self.q.supprimer_compte(self.courant)

    def test_virement_vers_le_meme_compte_refuse_par_la_base(self):
        with self.assertRaises(sqlite3.IntegrityError):
            self.q.enregistrer_operation(None, date(2025, 5, 1), "Virement", 1000, None, self.courant, None,
                                         self.courant, None)

    def test_export_excel(self):
        self.q.enregistrer_operation(None, date(2025, 5, 7), "Préventes", 1300, None, self.courant, self.ventes, None, None)
        self.q.enregistrer_operation(None, date(2026, 2, 2), "Virement livret", 1000, None, self.courant, None,
                                     self.livret, None)
        chemin = os.path.join(self.dossier.name, "bilan.xlsx")
        exporter_bilan(chemin, "BDE ODIN", Exercice(date(2025, 4, 1), date(2026, 3, 31)), self.q.get_tous_comptes(),
                       self.q.get_toutes_categories(), self.q.get_toutes_operations())

        wb = openpyxl.load_workbook(chemin)
        self.assertEqual(wb.sheetnames, ["Bilan financier", "Résumé financier"])
        journaux, resume = wb.worksheets
        self.assertEqual(journaux["B2"].value, "Bilan financier 2026 · BDE ODIN")
        self.assertEqual(journaux["B6"].value, "512 - Compte courant")
        self.assertEqual(journaux["C10"].value, "Préventes")
        self.assertEqual(journaux["C11"].value, "Sous-total")
        self.assertEqual(journaux["D11"].value, "=SUM(D10:D10)")
        self.assertEqual(journaux["D12"].value, "=D9+D11-G11")
        self.assertEqual(resume["B6"].value, "Compte courant")
        self.assertTrue(str(resume["B7"].value).startswith("='Bilan financier'!"))
        textes = {c.value for ligne in resume.iter_rows() for c in ligne if isinstance(c.value, str)}
        self.assertTrue({"dont adhésions", "Signature du président sortant", "Signature du trésorier entrant"} <= textes)


if __name__ == "__main__":
    unittest.main()
