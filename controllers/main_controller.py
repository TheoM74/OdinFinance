from views.main_window import MainWindow
from database.queries import TresoQueries
from PyQt6.QtWidgets import QTableWidgetItem, QMessageBox
from PyQt6.QtCore import QDate

class MainController:
    def __init__(self, view: MainWindow, queries: TresoQueries):
        self.view = view
        self.queries = queries
        self.transaction_en_cours_id = None  # None = Ajout, int = ID en cours de modification

        # Chargement initial
        self.charger_combobox()
        self.charger_donnees_completes()

        # Connexions des signaux
        self.view.btn_ajouter.clicked.connect(self.enregistrer_transaction)
        self.view.btn_annuler_edition.clicked.connect(self.reinitialiser_formulaire)
        
        self.view.btn_add_cat.clicked.connect(self.ouvrir_ajout_categorie)
        self.view.btn_del_cat.clicked.connect(self.supprimer_categorie)
        
        self.view.btn_add_tier.clicked.connect(self.ouvrir_ajout_tier)
        self.view.btn_del_tier.clicked.connect(self.supprimer_tier)

        self.view.btn_supprimer_trans.clicked.connect(self.supprimer_transaction_selectionnee)
        self.view.table_transactions.cellClicked.connect(self.charger_transaction_dans_formulaire)
        
        self.view.btn_export_excel.clicked.connect(self.exporter_bilan_excel)

    def charger_combobox(self):
        self.view.cat_combo.clear()
        self.categories_data = self.queries.get_toutes_categories()
        for cat_id, nom, type_flux in self.categories_data:
            self.view.cat_combo.addItem(f"[{type_flux}] {nom}", cat_id)

        self.view.tier_combo.clear()
        self.view.tier_combo.addItem("-- Aucun --", None)
        self.tiers_data = self.queries.get_tous_tiers()
        for tier_id, nom, type_tier in self.tiers_data:
            self.view.tier_combo.addItem(f"{nom} ({type_tier})", tier_id)

    def charger_donnees_completes(self):
        self.charger_tableau_transactions()
        self.charger_tableau_bilan()

    def charger_tableau_transactions(self):
        transactions = self.queries.get_toutes_transactions()
        self.view.table_transactions.setRowCount(len(transactions))

        # Cache pour retrouver l'ID complet lors du clic sur une ligne
        self.transactions_cache = transactions

        for row_idx, trans in enumerate(transactions):
            date_str = trans[1]
            desc = trans[2]
            montant = f"{trans[3]:.2f} €"
            ref = trans[4] if trans[4] else "-"
            cat = f"{trans[5]} ({trans[6]})"
            tier = trans[7] if trans[7] else "-"

            values = [date_str, desc, montant, cat, tier, ref]
            for col_idx, val in enumerate(values):
                item = QTableWidgetItem(val)
                self.view.table_transactions.setItem(row_idx, col_idx, item)

    def charger_tableau_bilan(self):
        totaux = self.queries.get_totaux_par_categorie()
        self.view.table_bilan.setRowCount(len(totaux))

        total_recettes = 0.0
        total_depenses = 0.0

        for row_idx, (nom_cat, type_flux, somme) in enumerate(totaux):
            if type_flux == "RECETTE":
                total_recettes += somme
            else:
                total_depenses += somme

            self.view.table_bilan.setItem(row_idx, 0, QTableWidgetItem(nom_cat))
            self.view.table_bilan.setItem(row_idx, 1, QTableWidgetItem(type_flux))
            self.view.table_bilan.setItem(row_idx, 2, QTableWidgetItem(f"{somme:.2f} €"))

        solde_net = total_recettes - total_depenses
        self.view.label_solde.setText(
            f"💰 Recettes : {total_recettes:.2f} €  |  "
            f"📉 Dépenses : {total_depenses:.2f} €  ||  "
            f"Solde Net : {solde_net:.2f} €"
        )

    def charger_transaction_dans_formulaire(self, row, col):
        """Remplit le formulaire lorsqu'on clique sur une ligne."""
        if row >= len(self.transactions_cache):
            return
        
        trans = self.transactions_cache[row]
        self.transaction_en_cours_id = trans[0]

        date_parts = list(map(int, trans[1].split("-")))
        self.view.date_input.setDate(QDate(date_parts[0], date_parts[1], date_parts[2]))
        self.view.desc_input.setText(trans[2])
        self.view.amount_input.setText(str(trans[3]))
        self.view.ref_input.setText(trans[4] if trans[4] else "")

        index_cat = self.view.cat_combo.findData(trans[8])
        if index_cat >= 0:
            self.view.cat_combo.setCurrentIndex(index_cat)

        if trans[9]:
            index_tier = self.view.tier_combo.findData(trans[9])
            if index_tier >= 0:
                self.view.tier_combo.setCurrentIndex(index_tier)
        else:
            self.view.tier_combo.setCurrentIndex(0)

        # Mode Modification UI
        self.view.form_group.setTitle(f"Modifier la transaction #{trans[0]}")
        self.view.btn_ajouter.setText("Mettre à jour")
        self.view.btn_ajouter.setStyleSheet("background-color: #f0ad4e; color: white; font-weight: bold; padding: 6px;")
        self.view.btn_annuler_edition.setVisible(True)

    def reinitialiser_formulaire(self):
        """Remet le formulaire en mode Ajout normal."""
        self.transaction_en_cours_id = None
        self.view.date_input.setDate(QDate.currentDate())
        self.view.desc_input.clear()
        self.view.amount_input.clear()
        self.view.ref_input.clear()
        self.view.tier_combo.setCurrentIndex(0)
        
        self.view.form_group.setTitle("Ajouter une transaction")
        self.view.btn_ajouter.setText("Enregistrer la transaction")
        self.view.btn_ajouter.setStyleSheet("background-color: #2b5b84; color: white; font-weight: bold; padding: 6px;")
        self.view.btn_annuler_edition.setVisible(False)

    def enregistrer_transaction(self):
        date_str = self.view.date_input.date().toString("yyyy-MM-dd")
        description = self.view.desc_input.text().strip()
        montant_str = self.view.amount_input.text().strip().replace(",", ".")
        ref_facture = self.view.ref_input.text().strip()
        
        categorie_id = self.view.cat_combo.currentData()
        tier_id = self.view.tier_combo.currentData()

        if not description:
            self.view.afficher_message("Erreur", "Veuillez saisir une description.", is_erreur=True)
            return

        try:
            montant = float(montant_str)
            if montant < 0:
                raise ValueError()
        except ValueError:
            self.view.afficher_message("Erreur", "Le montant doit être un nombre positif valide.", is_erreur=True)
            return

        try:
            if self.transaction_en_cours_id is None:
                self.queries.ajouter_transaction(
                    date=date_str, description=description, montant=montant,
                    ref_facture=ref_facture if ref_facture else None,
                    categorie_id=categorie_id, tier_id=tier_id
                )
                msg = "Transaction enregistrée avec succès !"
            else:
                self.queries.modifier_transaction(
                    trans_id=self.transaction_en_cours_id,
                    date=date_str, description=description, montant=montant,
                    ref_facture=ref_facture if ref_facture else None,
                    categorie_id=categorie_id, tier_id=tier_id
                )
                msg = f"Transaction mise à jour avec succès !"
                self.reinitialiser_formulaire()

            self.view.desc_input.clear()
            self.view.amount_input.clear()
            self.view.ref_input.clear()
            
            self.charger_donnees_completes()
            self.view.afficher_message("Succès", msg)
        except Exception as e:
            self.view.afficher_message("Erreur BDD", f"Erreur : {str(e)}", is_erreur=True)

    def supprimer_transaction_selectionnee(self):
        selected_rows = self.view.table_transactions.selectionModel().selectedRows()
        if not selected_rows:
            self.view.afficher_message("Attention", "Veuillez sélectionner une transaction à supprimer dans le tableau.", is_erreur=True)
            return

        row = selected_rows[0].row()
        trans = self.transactions_cache[row]
        trans_id = trans[0]

        confirm = QMessageBox.question(
            self.view, "Confirmation", 
            f"Voulez-vous vraiment supprimer la transaction du {trans[1]} ({trans[2]} - {trans[3]}€) ?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if confirm == QMessageBox.StandardButton.Yes:
            try:
                self.queries.supprimer_transaction(trans_id)
                self.charger_donnees_completes()
                self.reinitialiser_formulaire()
                self.view.afficher_message("Succès", "Transaction supprimée.")
            except Exception as e:
                self.view.afficher_message("Erreur", f"Impossible de supprimer : {str(e)}", is_erreur=True)

    def ouvrir_ajout_categorie(self):
        nom, type_flux = self.view.demander_texte_et_type("Ajouter une catégorie", ["RECETTE", "DEPENSE"])
        if nom:
            try:
                self.queries.ajouter_categorie(nom, type_flux)
                self.charger_combobox()
                self.view.afficher_message("Succès", f"Catégorie '{nom}' ajoutée !")
            except Exception as e:
                self.view.afficher_message("Erreur", f"Impossible d'ajouter : {str(e)}", is_erreur=True)

    def supprimer_categorie(self):
        cat_id = self.view.cat_combo.currentData()
        if not cat_id:
            return
        
        cat_text = self.view.cat_combo.currentText()
        confirm = QMessageBox.question(
            self.view, "Confirmation", 
            f"Voulez-vous supprimer la catégorie '{cat_text}' ?\n(Échouera si des transactions l'utilisent)",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if confirm == QMessageBox.StandardButton.Yes:
            try:
                self.queries.supprimer_categorie(cat_id)
                self.charger_combobox()
                self.view.afficher_message("Succès", "Catégorie supprimée.")
            except Exception as e:
                self.view.afficher_message("Erreur", "Impossible : cette catégorie est sûrement utilisée dans des transactions.", is_erreur=True)

    def ouvrir_ajout_tier(self):
        nom, type_tier = self.view.demander_texte_et_type("Ajouter un tier", ["MORAL", "PHYSIQUE"])
        if nom:
            try:
                self.queries.ajouter_tier(nom, type_tier)
                self.charger_combobox()
                self.view.afficher_message("Succès", f"Tier '{nom}' ajouté !")
            except Exception as e:
                self.view.afficher_message("Erreur", f"Impossible d'ajouter : {str(e)}", is_erreur=True)

    def supprimer_tier(self):
        tier_id = self.view.tier_combo.currentData()
        if not tier_id:
            self.view.afficher_message("Attention", "Veuillez sélectionner un tier valide à supprimer.", is_erreur=True)
            return
        
        tier_text = self.view.tier_combo.currentText()
        confirm = QMessageBox.question(
            self.view, "Confirmation", 
            f"Voulez-vous supprimer le tier '{tier_text}' ?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )

        if confirm == QMessageBox.StandardButton.Yes:
            try:
                self.queries.supprimer_tier(tier_id)
                self.charger_combobox()
                self.view.afficher_message("Succès", "Tier supprimé.")
            except Exception as e:
                self.view.afficher_message("Erreur", f"Impossible : ce tier est lié à des transactions.", is_erreur=True)

    def exporter_bilan_excel(self):
        from PyQt6.QtWidgets import QFileDialog
        import openpyxl  # type: ignore[reportMissingModuleSource]
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore[reportMissingModuleSource]
        from datetime import datetime

        file_path, _ = QFileDialog.getSaveFileName(
            self.view, "Enregistrer le Bilan AG", 
            f"Bilan_Financier_ODIN_{datetime.now().year}.xlsx", 
            "Fichiers Excel (*.xlsx)"
        )

        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Bilan AG"
            ws.views.sheetView[0].showGridLines = True

            font_titre = Font(name="Calibri", size=16, bold=True, color="1F497D")
            font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            fill_header = PatternFill(start_color="2B5B84", end_color="2B5B84", fill_type="solid")
            border_thin = Border(
                left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
            )

            ws.append(["RAPPORT FINANCIER ANNUEL - ASSOCIATION ODIN"])
            ws.cell(row=1, column=1).font = font_titre
            ws.append([f"Généré le : {datetime.now().strftime('%d/%m/%Y à %H:%M')}"])
            ws.append([])

            headers = ["Catégorie", "Type de Flux", "Total Cumulé (€)"]
            ws.append(headers)
            for col_num in range(1, 4):
                cell = ws.cell(row=4, column=col_num)
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = Alignment(horizontal="center", vertical="center")

            totaux = self.queries.get_totaux_par_categorie()
            total_recettes = 0.0
            total_depenses = 0.0

            for row_data in totaux:
                nom_cat, type_flux, somme = row_data
                if type_flux == "RECETTE":
                    total_recettes += somme
                else:
                    total_depenses += somme
                
                ws.append([nom_cat, type_flux, somme])
                current_row = ws.max_row
                ws.cell(row=current_row, column=3).number_format = '#,##0.00 "€"'
                for col_num in range(1, 4):
                    ws.cell(row=current_row, column=col_num).border = border_thin

            ws.append([])
            solde_net = total_recettes - total_depenses
            
            ws.append(["Total Recettes", "", total_recettes])
            ws.cell(row=ws.max_row, column=3).number_format = '#,##0.00 "€"'
            ws.cell(row=ws.max_row, column=3).font = Font(bold=True, color="008000")

            ws.append(["Total Dépenses", "", total_depenses])
            ws.cell(row=ws.max_row, column=3).number_format = '#,##0.00 "€"'
            ws.cell(row=ws.max_row, column=3).font = Font(bold=True, color="FF0000")

            ws.append(["Solde Net en Caisse", "", solde_net])
            ws.cell(row=ws.max_row, column=3).number_format = '#,##0.00 "€"'
            ws.cell(row=ws.max_row, column=3).font = Font(bold=True, color="1F497D")

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

            wb.save(file_path)
            self.view.afficher_message("Succès", f"Bilan AG exporté avec succès :\n{file_path}")

        except Exception as e:
            self.view.afficher_message("Erreur", f"Erreur lors de l'export Excel : {str(e)}", is_erreur=True)